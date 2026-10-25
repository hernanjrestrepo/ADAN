"""Extracción de texto de la documentación que el cliente carga (plan de negocio, estados financieros,
presentaciones, hojas de cálculo, actas...). Todo se procesa en memoria: el archivo original no se guarda,
solo su texto, que pasa a ser un Documento del cliente y entra a su memoria empresarial (EMS).
"""
from __future__ import annotations

import csv
import io
import zipfile
from dataclasses import dataclass

MAX_FILE_BYTES = 15 * 1024 * 1024
MAX_MEDIA_BYTES = 50 * 1024 * 1024  # grabaciones de audio y video
MAX_TEXT_CHARS = 200_000
MAX_XLSX_CELLS = 60_000
# Un .docx/.xlsx/.pptx es un ZIP: se rechaza si descomprimido pasa de esto (bomba de compresión)
MAX_UNZIPPED_BYTES = 120 * 1024 * 1024

SUPPORTED = {
    ".pdf": "PDF", ".docx": "Word", ".xlsx": "Excel", ".pptx": "PowerPoint",
    ".txt": "Texto", ".md": "Texto", ".csv": "CSV", ".json": "Texto",
    # Imágenes y capturas (OCR), audio y video (transcripción): ver media.py
    ".png": "Imagen", ".jpg": "Imagen", ".jpeg": "Imagen", ".webp": "Imagen", ".gif": "Imagen",
    ".mp3": "Audio", ".m4a": "Audio", ".wav": "Audio", ".ogg": "Audio", ".opus": "Audio", ".aac": "Audio",
    ".flac": "Audio", ".amr": "Audio",
    ".mp4": "Video", ".mov": "Video", ".webm": "Video", ".mkv": "Video", ".avi": "Video", ".3gp": "Video",
}
MEDIA_KINDS = ("Imagen", "Audio", "Video")


class ExtractionError(ValueError):
    """El archivo no se puede leer: formato no admitido, dañado, protegido o demasiado grande."""


@dataclass
class Extracted:
    filename: str
    kind: str
    text: str
    pages: int
    truncated: bool
    # Lo que falta por hacer fuera de aquí: "ocr" (imagen o PDF escaneado) o "transcribir" (audio o video)
    pending: str = ""
    data: bytes = b""


def extension(filename: str) -> str:
    name = filename.lower().strip()
    return name[name.rfind("."):] if "." in name else ""


def extract(filename: str, data: bytes) -> Extracted:
    ext = extension(filename)
    if ext not in SUPPORTED:
        raise ExtractionError(f"{filename}: formato no admitido. Sube PDF, Word, Excel, PowerPoint, CSV, texto, "
                              "imágenes, audio o video.")
    kind = SUPPORTED[ext]
    limit = MAX_MEDIA_BYTES if kind in ("Audio", "Video") else MAX_FILE_BYTES
    if len(data) > limit:
        raise ExtractionError(f"{filename}: pesa más de {limit // (1024 * 1024)} MB.")
    if not data:
        raise ExtractionError(f"{filename}: el archivo está vacío.")
    if kind in MEDIA_KINDS:
        return Extracted(filename=filename, kind=kind, text="", pages=1, truncated=False,
                         pending="ocr" if kind == "Imagen" else "transcribir", data=data)
    if ext in (".docx", ".xlsx", ".pptx"):
        _check_zip(filename, data)
    try:
        text, pages = _EXTRACTORS[ext](data)
    except ExtractionError:
        raise
    except Exception as exc:  # archivo dañado o protegido con clave
        raise ExtractionError(f"{filename}: no se pudo leer ({type(exc).__name__}). ¿Está dañado o protegido?") from exc
    text = _clean(text)
    if ext == ".pdf" and len(text.strip()) < 40 * max(pages, 1):
        # PDF escaneado: sus páginas son imágenes; se leen con OCR
        return Extracted(filename=filename, kind="PDF", text=text, pages=pages, truncated=False, pending="ocr", data=data)
    if len(text) < 20:
        raise ExtractionError(f"{filename}: no tiene texto legible (¿es una imagen escaneada?).")
    truncated = len(text) > MAX_TEXT_CHARS
    return Extracted(filename=filename, kind=SUPPORTED[ext], text=text[:MAX_TEXT_CHARS], pages=pages, truncated=truncated)


def _check_zip(filename: str, data: bytes) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            if sum(i.file_size for i in z.infolist()) > MAX_UNZIPPED_BYTES:
                raise ExtractionError(f"{filename}: el contenido descomprimido es demasiado grande.")
    except zipfile.BadZipFile as exc:
        raise ExtractionError(f"{filename}: no es un archivo de Office válido.") from exc


def _clean(text: str) -> str:
    lines = [" ".join(line.split()) for line in text.replace("\x00", "").splitlines()]
    out, blank = [], 0
    for line in lines:
        blank = blank + 1 if not line else 0
        if blank <= 1:
            out.append(line)
    return "\n".join(out).strip()


def _pdf(data: bytes) -> tuple[str, int]:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted:
        raise ExtractionError("El PDF está protegido con clave.")
    parts = []
    for i, page in enumerate(reader.pages, 1):
        parts.append(f"[Página {i}]\n{page.extract_text() or ''}")
        if sum(len(p) for p in parts) > MAX_TEXT_CHARS:
            break
    return "\n\n".join(parts), len(reader.pages)


def _docx(data: bytes) -> tuple[str, int]:
    from docx import Document

    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(c.text.strip() for c in row.cells))
    return "\n".join(parts), 1


def _xlsx(data: bytes) -> tuple[str, int]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    parts, cells = [], 0
    for ws in wb.worksheets:
        parts.append(f"[Hoja: {ws.title}]")
        for row in ws.iter_rows(values_only=True):
            values = ["" if v is None else str(v) for v in row]
            if any(values):
                parts.append(" | ".join(values).rstrip(" |"))
            cells += len(values)
            if cells > MAX_XLSX_CELLS:
                parts.append("[...hoja recortada: demasiadas celdas]")
                break
    count = len(wb.worksheets)
    wb.close()
    return "\n".join(parts), count


def _pptx(data: bytes) -> tuple[str, int]:
    from pptx import Presentation

    prs = Presentation(io.BytesIO(data))
    parts = []
    for i, slide in enumerate(prs.slides, 1):
        parts.append(f"[Diapositiva {i}]")
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False) and shape.text_frame.text.strip():
                parts.append(shape.text_frame.text)
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    parts.append(" | ".join(c.text.strip() for c in row.cells))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            parts.append("Notas: " + slide.notes_slide.notes_text_frame.text)
    return "\n".join(parts), len(prs.slides)


def _text(data: bytes) -> tuple[str, int]:
    for encoding in ("utf-8-sig", "latin-1"):
        try:
            return data.decode(encoding), 1
        except UnicodeDecodeError:
            continue
    raise ExtractionError("No se reconoce la codificación del texto.")


def _csv(data: bytes) -> tuple[str, int]:
    text, _ = _text(data)
    rows = list(csv.reader(io.StringIO(text)))
    return "\n".join(" | ".join(r) for r in rows[:5000]), 1


_EXTRACTORS = {".pdf": _pdf, ".docx": _docx, ".xlsx": _xlsx, ".pptx": _pptx, ".txt": _text, ".md": _text,
               ".json": _text, ".csv": _csv}
