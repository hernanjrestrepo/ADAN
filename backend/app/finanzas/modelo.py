"""Motor del flujo de caja: con los supuestos (que propone el CFO de ADAN y el cliente puede corregir), calcula
mes a mes ingresos, costos, EBITDA, impuestos, inversión, financiación y caja. Las cuentas las hace el código,
no el modelo de lenguaje: son exactas y repetibles.
"""
from __future__ import annotations

import io
import math

MONTHS = 24
SCENARIOS = {"pesimista": {"units": 0.6, "growth": 0.5, "price": 0.95},
             "base": {"units": 1.0, "growth": 1.0, "price": 1.0},
             "optimista": {"units": 1.3, "growth": 1.25, "price": 1.0}}


def _num(value, default: float = 0.0) -> float:
    try:
        n = float(value)
        return n if math.isfinite(n) else default
    except (TypeError, ValueError):
        return default


def normalize(a: dict) -> dict:
    """Supuestos con tipos y límites sanos, venga de donde venga (modelo o edición del cliente)."""
    lines = lambda key, fields: [  # noqa: E731
        {**{f: (x.get(f) if f in ("concepto", "fuente", "base", "tipo") else _num(x.get(f))) for f in fields}}
        for x in (a.get(key) or []) if isinstance(x, dict) and str(x.get("concepto") or x.get("fuente") or "").strip()
    ]
    return {
        "moneda": str(a.get("moneda") or "COP")[:5],
        "precio_unitario": max(0.0, _num(a.get("precio_unitario"))),
        "unidad": str(a.get("unidad") or "venta")[:40],
        "mes_inicio_ventas": int(min(MONTHS, max(1, _num(a.get("mes_inicio_ventas"), 1)))),
        "unidades_mes_1": max(0.0, _num(a.get("unidades_mes_1"))),
        "crecimiento_mensual_pct": min(100.0, max(-50.0, _num(a.get("crecimiento_mensual_pct")))),
        "unidades_max_mes": max(0.0, _num(a.get("unidades_max_mes"), 1e12)) or 1e12,
        "costo_variable_pct": min(100.0, max(0.0, _num(a.get("costo_variable_pct")))),
        "comision_pagos_pct": min(20.0, max(0.0, _num(a.get("comision_pagos_pct")))),
        "marketing_mensual": max(0.0, _num(a.get("marketing_mensual"))),
        "impuesto_renta_pct": min(60.0, max(0.0, _num(a.get("impuesto_renta_pct")))),
        "capital_inicial": max(0.0, _num(a.get("capital_inicial"))),
        "gastos_fijos": [{**g, "monto_mensual": max(0.0, g["monto_mensual"]),
                          "desde_mes": int(min(MONTHS, max(1, g["desde_mes"] or 1)))}
                         for g in lines("gastos_fijos", ("concepto", "monto_mensual", "desde_mes", "tipo", "base"))],
        "capex": [{**c, "monto": max(0.0, c["monto"]), "mes": int(min(MONTHS, max(1, c["mes"] or 1)))}
                  for c in lines("capex", ("concepto", "monto", "mes", "base"))],
        "financiacion": [{**f, "monto": max(0.0, f["monto"]), "mes": int(min(MONTHS, max(1, f["mes"] or 1)))}
                         for f in lines("financiacion", ("fuente", "monto", "mes", "base"))],
        "bases": {k: str(v)[:300] for k, v in (a.get("bases") or {}).items()} if isinstance(a.get("bases"), dict) else {},
    }


def compute(raw: dict, scenario: str = "base") -> dict:
    a = normalize(raw)
    k = SCENARIOS.get(scenario, SCENARIOS["base"])
    price = a["precio_unitario"] * k["price"]
    growth = a["crecimiento_mensual_pct"] * k["growth"] / 100
    start = a["mes_inicio_ventas"]
    first_units = a["unidades_mes_1"] * k["units"]
    units = 0.0
    cash = a["capital_inicial"]
    cash_without_funding = a["capital_inicial"]
    rows = []
    for m in range(1, MONTHS + 1):
        # Las ventas arrancan en `mes_inicio_ventas` con `unidades_mes_1` y desde ahí crecen
        if m == start:
            units = min(a["unidades_max_mes"], first_units)
        elif m > start:
            units = min(a["unidades_max_mes"], units * (1 + growth))
        revenue = units * price
        variable = revenue * (a["costo_variable_pct"] + a["comision_pagos_pct"]) / 100
        fixed = sum(g["monto_mensual"] for g in a["gastos_fijos"] if g["desde_mes"] <= m)
        ebitda = revenue - variable - fixed - a["marketing_mensual"]
        taxes = max(0.0, ebitda) * a["impuesto_renta_pct"] / 100
        capex = sum(c["monto"] for c in a["capex"] if c["mes"] == m)
        funding = sum(f["monto"] for f in a["financiacion"] if f["mes"] == m)
        net = ebitda - taxes - capex
        cash += net + funding
        cash_without_funding += net
        rows.append({"mes": m, "unidades": round(units, 1), "ingresos": round(revenue), "costo_variable": round(variable),
                     "margen_contribucion": round(revenue - variable), "gastos_fijos": round(fixed),
                     "marketing": round(a["marketing_mensual"]), "ebitda": round(ebitda), "impuestos": round(taxes),
                     "capex": round(capex), "financiacion": round(funding), "flujo_neto": round(net), "caja": round(cash),
                     "_sin_fin": cash_without_funding})
    min_row = min(rows, key=lambda r: r["caja"])
    lowest_without_funding = min(r["_sin_fin"] for r in rows)
    burn = sum(r["gastos_fijos"] + r["marketing"] for r in rows[:3]) / 3
    breakeven = next((r["mes"] for r in rows if r["ebitda"] >= 0 and r["ingresos"] > 0), None)
    year = lambda key, y: sum(r[key] for r in rows[(y - 1) * 12: y * 12])  # noqa: E731
    revenue_total = sum(r["ingresos"] for r in rows)
    kpis = {
        "punto_equilibrio_mes": breakeven,
        "caja_minima": min_row["caja"], "caja_minima_mes": min_row["mes"],
        # Lo que hace falta para no quedar en rojo sin financiación, con un colchón de 3 meses de gastos fijos
        "capital_necesario": round(max(0.0, -lowest_without_funding) + 3 * burn),
        "margen_contribucion_pct": round(100 - a["costo_variable_pct"] - a["comision_pagos_pct"], 1),
        "ingresos_ano_1": year("ingresos", 1), "ingresos_ano_2": year("ingresos", 2),
        "ebitda_ano_1": year("ebitda", 1), "ebitda_ano_2": year("ebitda", 2),
        "caja_final": rows[-1]["caja"], "se_queda_sin_caja": min_row["caja"] < 0,
        "mes_sin_caja": next((r["mes"] for r in rows if r["caja"] < 0), None),
        "gasto_fijo_mensual_inicial": round(burn),
        "ingresos_24_meses": revenue_total,
    }
    for r in rows:
        r.pop("_sin_fin")
    return {"scenario": scenario, "rows": rows, "kpis": kpis}


def full(raw: dict) -> dict:
    """Supuestos normalizados, el escenario base mes a mes y los indicadores de los tres escenarios."""
    base = compute(raw, "base")
    return {"assumptions": normalize(raw), "base": base,
            "scenarios": {name: compute(raw, name)["kpis"] for name in SCENARIOS}}


def to_xlsx(model: dict, company: str) -> bytes:
    """El modelo en Excel, para el contador o el inversionista."""
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "Flujo de caja"
    cols = ["mes", "unidades", "ingresos", "costo_variable", "margen_contribucion", "gastos_fijos", "marketing", "ebitda",
            "impuestos", "capex", "financiacion", "flujo_neto", "caja"]
    ws.append([f"Flujo de caja — {company} ({model['assumptions']['moneda']}), escenario base"])
    ws.append([c.replace("_", " ").capitalize() for c in cols])
    for cell in ws[2]:
        cell.font = Font(bold=True)
    for r in model["base"]["rows"]:
        ws.append([r[c] for c in cols])
    s = wb.create_sheet("Supuestos")
    a = model["assumptions"]
    for key in ("moneda", "precio_unitario", "unidad", "mes_inicio_ventas", "unidades_mes_1", "crecimiento_mensual_pct", "unidades_max_mes",
                "costo_variable_pct", "comision_pagos_pct", "marketing_mensual", "impuesto_renta_pct", "capital_inicial"):
        s.append([key.replace("_", " "), a[key], a["bases"].get(key, "")])
    s.append([])
    s.append(["Gastos fijos", "Mensual", "Desde el mes", "Tipo", "Base"])
    for g in a["gastos_fijos"]:
        s.append([g["concepto"], g["monto_mensual"], g["desde_mes"], g.get("tipo") or "", g.get("base") or ""])
    k = wb.create_sheet("Escenarios")
    k.append(["Indicador", *model["scenarios"].keys()])
    for key in model["scenarios"]["base"]:
        k.append([key.replace("_", " "), *[model["scenarios"][sc][key] for sc in model["scenarios"]]])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
