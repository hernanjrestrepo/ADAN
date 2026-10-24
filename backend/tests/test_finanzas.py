"""Finanzas: el flujo de caja lo calcula el código (cifras exactas); el organigrama híbrido entra al flujo."""
import io
import json

from app.ai.base import LLMAdapter, LLMResponse
from app.finanzas import modelo
from tests.test_discovery import _setup, _use

BASE = {"moneda": "COP", "precio_unitario": 100000, "unidad": "suscripción", "unidades_mes_1": 10,
        "crecimiento_mensual_pct": 10, "unidades_max_mes": 100, "costo_variable_pct": 30, "comision_pagos_pct": 3,
        "marketing_mensual": 500000, "impuesto_renta_pct": 35, "capital_inicial": 5000000,
        "gastos_fijos": [{"concepto": "Arriendo", "monto_mensual": 1000000, "desde_mes": 1},
                         {"concepto": "Vendedor", "monto_mensual": 2000000, "desde_mes": 4}],
        "capex": [{"concepto": "Horno", "monto": 3000000, "mes": 1}],
        "financiacion": [{"fuente": "Socio", "monto": 10000000, "mes": 6}]}


def test_the_cash_flow_is_computed_exactly():
    m = modelo.compute(BASE)
    first, fourth = m["rows"][0], m["rows"][3]
    assert first["ingresos"] == 1_000_000 and first["costo_variable"] == 330_000
    assert first["ebitda"] == 1_000_000 - 330_000 - 1_000_000 - 500_000
    assert first["caja"] == 5_000_000 - 830_000 - 3_000_000  # sin impuestos con EBITDA negativo, menos el horno
    assert fourth["gastos_fijos"] == 3_000_000  # el vendedor entra en el mes 4
    assert round(fourth["unidades"], 1) == 13.3  # 10 × 1,1³
    assert m["rows"][5]["financiacion"] == 10_000_000
    assert max(r["unidades"] for r in m["rows"]) <= 100  # techo de capacidad


def test_kpis_and_scenarios():
    full = modelo.full(BASE)
    k = full["base"]["kpis"]
    assert k["margen_contribucion_pct"] == 67.0
    assert k["punto_equilibrio_mes"] is not None and k["capital_necesario"] > 0
    assert k["mes_sin_caja"] == next((r["mes"] for r in full["base"]["rows"] if r["caja"] < 0), None)
    assert full["scenarios"]["pesimista"]["ingresos_24_meses"] < k["ingresos_24_meses"] < full["scenarios"]["optimista"]["ingresos_24_meses"]


def test_bad_inputs_are_sanitized():
    a = modelo.normalize({"precio_unitario": "abc", "crecimiento_mensual_pct": 900, "gastos_fijos": [
        {"concepto": "", "monto_mensual": 1}, {"concepto": "Luz", "monto_mensual": -5, "desde_mes": 99}]})
    assert a["precio_unitario"] == 0 and a["crecimiento_mensual_pct"] == 100
    assert a["gastos_fijos"] == [{"concepto": "Luz", "monto_mensual": 0.0, "desde_mes": 24, "tipo": None, "base": None}]


ORG = {"moneda": "cop", "tasa_usd": 4000, "resumen": "Equipo pequeño con agentes.", "roles": [
    {"area": "Dirección", "cargo": "Gerente", "tipo": "humano", "servicio": "", "responsabilidades": ["Dirigir"],
     "desde_mes": 1, "minutos_mes": 0, "costo_mensual": 5000000, "costo_humano_equivalente": 5000000, "por_que": "Responsabilidad legal"},
    {"area": "Comercial", "cargo": "Ventas y marketing", "tipo": "servicio_paradixe", "servicio": "EVA", "responsabilidades": ["Campañas"],
     "desde_mes": 2, "minutos_mes": 3000, "costo_mensual": 999, "costo_humano_equivalente": 4000000, "por_que": "Campañas 24/7"}]}
ASSUMPTIONS = {"moneda": "USD", "unidad": "suscripción", "precio_unitario": 100000, "mes_inicio_ventas": 1, "unidades_mes_1": 0,
               "crecimiento_mensual_pct": 15, "unidades_max_mes": 500, "costo_variable_pct": 35, "comision_pagos_pct": 3,
               "marketing_mensual": 800000, "impuesto_renta_pct": 35, "capital_inicial": 0,
               "gastos_fijos_adicionales": [{"concepto": "Arriendo", "monto": 1200000, "mes": 1, "base": "benchmark: arriendo local"}],
               "capex": [], "bases": {"precio_unitario": "supuesto: precio de la competencia"}, "advertencias": ["Precio sin validar"]}
REPORT = {"titulo": "Informe", "respuesta": "No está lista: falta validar el precio.",
          "scqa": {"situacion": "s", "complicacion": "c", "pregunta": "p", "respuesta": "r"},
          "argumentos": [{"titulo": "Mercado", "evidencia": ["DANE 2025"]}],
          "lean_canvas": {k: k for k in ("problema", "segmentos", "propuesta_unica", "solucion", "canales", "ingresos", "costos",
                                         "metricas_clave", "ventaja_injusta")},
          "benchmark": [{"competidor": "Despegar", "fortalezas": "f", "debilidades": "d", "como_ganarle": "g"}],
          "riesgos": [{"riesgo": "Caja", "probabilidad": "alta", "impacto": "alto", "mitigacion": "m"}],
          "six_sigma": {"ctq": ["Puntualidad"], "definir": "d", "medir": "m", "analizar": "a", "mejorar": "me", "controlar": "c"},
          "plan_90_dias": [{"semanas": "1-2", "accion": "Entrevistas", "responsable": "Fundador", "entregable": "10 entrevistas"}],
          "recomendacion": "Validar antes de invertir."}


class CfoLLM(LLMAdapter):
    def __init__(self):
        self.prompts = []

    async def chat_json(self, messages, schema, model=None, temperature=0.3, max_tokens=2048):
        self.prompts.append(messages[-1].content)
        props = schema["properties"]
        data = ORG if "roles" in props else ASSUMPTIONS if "precio_unitario" in props else REPORT
        return LLMResponse(content=json.dumps(data), model="fake", parsed=data)

    async def chat(self, *a, **k):
        raise AssertionError

    async def generate(self, *a, **k):
        raise AssertionError

    async def health_check(self):
        return True

    def list_models(self):
        return []


def test_org_chart_feeds_the_cash_flow_and_the_report(client):
    cid, headers = _setup(client, "finanzas@example.com")
    llm = _use(client, CfoLLM())
    org = client.post(f"/api/v1/finanzas/{cid}/organigrama", headers=headers).json()
    # EVA: 3.000 min × USD 0,10 × 4.000 = COP 1.200.000 (precio real de Paradixe, no el estimado del modelo);
    # la membresía de USD 100 (COP 400.000) entra una vez, desde el primer mes con servicios de Paradixe
    eva = org["roles"][1]
    assert eva["costo_mensual"] == 1_200_000 and "USD 0.10 por minuto × 3,000 min" in eva["precio"]
    membership = org["roles"][-1]
    assert membership["servicio"] == "Membresía Paradixe" and membership["costo_mensual"] == 400_000 and membership["desde_mes"] == 2
    assert org["costo_mensual_total"] == 6_600_000 and org["ahorro_mensual"] == 2_800_000
    assert "EVA" in llm.prompts[0] and "BPO Paradixe" in llm.prompts[0]

    model = client.post(f"/api/v1/finanzas/{cid}/modelo", headers=headers).json()
    concepts = [g["concepto"] for g in model["assumptions"]["gastos_fijos"]]
    assert "Gerente (persona)" in concepts and "Ventas y marketing (EVA)" in concepts and "Arriendo" in concepts
    assert model["base"]["rows"][0]["gastos_fijos"] == 5_000_000 + 1_200_000  # EVA y la membresía entran en el mes 2
    assert model["base"]["rows"][1]["gastos_fijos"] == 6_600_000 + 1_200_000
    assert model["advertencias"] == ["Precio sin validar"] and "organigrama ya cuesta" in llm.prompts[-1]
    # Una sola moneda: la del organigrama manda aunque el CFO proponga otra
    assert org["moneda"] == "COP" and model["assumptions"]["moneda"] == "COP" and "MONEDA OBLIGATORIA: COP" in llm.prompts[-1]

    edited = dict(model["assumptions"], precio_unitario=200000, unidades_mes_1=20)
    again = client.post(f"/api/v1/finanzas/{cid}/modelo/calcular", json={"assumptions": edited}, headers=headers).json()
    assert again["base"]["rows"][0]["ingresos"] == 4_000_000 and again["edited"] is True

    xlsx = client.get(f"/api/v1/finanzas/{cid}/modelo.xlsx", headers=headers)
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(xlsx.content))
    assert wb.sheetnames == ["Flujo de caja", "Supuestos", "Escenarios"]

    report = client.post(f"/api/v1/finanzas/{cid}/informe", headers=headers).json()
    assert report["respuesta"].startswith("No está lista") and report["kpis"]["margen_contribucion_pct"] == 62.0
    assert "punto de equilibrio en el mes" in llm.prompts[-1]
    docx = client.get(f"/api/v1/finanzas/{cid}/informe.docx", headers=headers)
    from docx import Document as Docx
    text = "\n".join(p.text for p in Docx(io.BytesIO(docx.content)).paragraphs)
    assert "Lean Canvas" in text and "Six Sigma" in text and "Plan de 90 días" in text


def test_finance_of_another_company_is_off_limits(client):
    cid, _ = _setup(client, "finduena@example.com")
    _, other = _setup(client, "finintrusa@example.com")
    _use(client, CfoLLM())
    assert client.post(f"/api/v1/finanzas/{cid}/organigrama", headers=other).status_code == 404
    assert client.get(f"/api/v1/finanzas/{cid}/modelo", headers=other).status_code == 404


def test_sales_start_in_a_given_month_and_then_grow():
    m = modelo.compute({**BASE, "mes_inicio_ventas": 4, "unidades_mes_1": 5, "crecimiento_mensual_pct": 20, "capex": [],
                        "financiacion": []})
    assert [r["unidades"] for r in m["rows"][:6]] == [0, 0, 0, 5, 6, 7.2]
    assert m["rows"][2]["ingresos"] == 0 and m["rows"][3]["ingresos"] == 500_000
