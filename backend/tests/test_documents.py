"""Documentación del cliente: ADÁN extrae el texto de PDF, Word, Excel y PowerPoint y aprende de él."""
import io

import pytest

from app.documents import extract as ex
from app.models.models import Document, Message
from app.nivel1.service import Nivel1Service
from tests.test_discovery import ScriptLLM, _setup, _use, topic


def _docx(text: str) -> bytes:
    from docx import Document as Docx

    d = Docx()
    d.add_paragraph(text)
    t = d.add_table(rows=1, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Ventas 2025", "COP 120 millones"
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def _xlsx() -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Flujo de caja"
    ws.append(["Mes", "Ingresos", "Gastos"])
    ws.append(["Enero", 1000000, 800000])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _pptx() -> bytes:
    from pptx import Presentation

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Nuestra propuesta de valor"
    slide.placeholders[1].text = "Pan fresco a las 6:30 a. m. por suscripción"
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _pdf(text: str) -> bytes:
    """PDF mínimo de una página con texto (sin dependencias para generarlo)."""
    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
    objs = [b"<< /Type /Catalog /Pages 2 0 R >>", b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
            b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
            b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    out, offsets = b"%PDF-1.4\n", []
    for i, body in enumerate(objs, 1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n".encode() + body + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode()
    out += b"".join(f"{o:010d} 00000 n \n".encode() for o in offsets)
    out += f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode()
    return out


def test_extracts_every_supported_format():
    assert "COP 120 millones" in ex.extract("plan.docx", _docx("Plan de negocio de la panadería")).text
    sheet = ex.extract("finanzas.xlsx", _xlsx())
    assert "[Hoja: Flujo de caja]" in sheet.text and "Enero | 1000000 | 800000" in sheet.text
    assert "suscripción" in ex.extract("pitch.pptx", _pptx()).text
    pdf = ex.extract("estudio.pdf", _pdf("Estudio de mercado de panaderias en Barranquilla"))
    assert "Barranquilla" in pdf.text and pdf.pages == 1
    assert ex.extract("notas.md", "# Notas\nClientes: familias del barrio".encode()).kind == "Texto"
    assert "familias | 120" in ex.extract("datos.csv", b"segmento,clientes\nfamilias,120").text


@pytest.mark.parametrize("name,data,message", [
    ("virus.exe", b"MZ....", "formato no admitido"),
    ("vacio.txt", b"", "vacío"),
    ("falso.docx", b"no soy un zip", "no es un archivo de Office"),
    ("corto.txt", b"hola", "no tiene texto legible"),
])
def test_rejects_what_cannot_be_read(name, data, message):
    with pytest.raises(ex.ExtractionError, match=message):
        ex.extract(name, data)


def test_rejects_files_that_are_too_big():
    with pytest.raises(ex.ExtractionError, match="pesa más"):
        ex.extract("enorme.txt", b"a" * (ex.MAX_FILE_BYTES + 1))


LEARNING = {
    "respuesta": "Leí tu plan: vendes pan de masa madre por suscripción. Lo más débil: no hay ventas reales. ¿Cuánto les cuesta hoy?",
    "resumen": "Panadería de barrio que vende pan de masa madre por suscripción semanal a familias.",
    "etapa": "idea",
    "hechos": [{"dato": "El plan proyecta 120 suscriptores", "fuente": "plan.docx", "base": "supuesto"}],
    "temas": [topic("problema", resumen="Pan fresco antes de las 7 a. m.", base="supuesto"),
              topic("afectados", estado="parcial", resumen="Familias del barrio")],
    "evidencia_sugerida": [{"afirmacion": "Estudio de mercado de panaderías en Barranquilla (2025)", "tipo": "external",
                            "fuente": "estudio.pdf"}],
    "vacios": ["No hay estructura de costos"],
    "riesgos": ["Un solo horno limita la producción"],
}


def test_upload_saves_documents_and_adan_learns_from_them(client, db_session):
    cid, headers = _setup(client, "docs@example.com")
    llm = _use(client, ScriptLLM([LEARNING]))
    files = [("files", ("plan.docx", _docx("Plan de negocio de la panadería Masa Madre"), "application/octet-stream")),
             ("files", ("estudio.pdf", _pdf("Estudio de mercado de panaderias en Barranquilla"), "application/pdf")),
             ("files", ("foto.exe", b"MZ", "application/octet-stream"))]
    resp = client.post(f"/api/v1/nivel1/{cid}/documents/upload", files=files, headers=headers)
    assert resp.status_code == 200, resp.text
    body = resp.json()
    assert [d["title"] for d in body["documents"]] == ["plan.docx", "estudio.pdf"]
    assert any("foto.exe" in e for e in body["errors"])
    # El modelo leyó los dos documentos completos
    corpus = llm.calls[0][1].content
    assert "=== plan.docx ===" in corpus and "Barranquilla" in corpus
    # El guion avanza con lo que dicen los documentos y la evidencia sugerida aparece
    assert body["discovery"]["done"] == 1
    assert body["discovery"]["evidence_suggestions"][0]["fuente"] == "estudio.pdf"
    assert body["learning"]["riesgos"] == ["Un solo horno limita la producción"]
    assert body["reply"].startswith("📎 Leí 2 documento(s)")
    stored = db_session.query(Document).filter_by(doc_type="client_upload").all()
    assert {d.origin for d in stored} == {"received_from_client"}
    assert db_session.query(Message).filter(Message.content.like("📎 Leí%")).count() == 1


def test_the_board_reads_what_adan_learned_from_the_documents(client, db_session):
    cid, headers = _setup(client, "docsboard@example.com")
    _use(client, ScriptLLM([LEARNING]))
    client.post(f"/api/v1/nivel1/{cid}/documents/upload", headers=headers,
                files=[("files", ("plan.docx", _docx("Plan de negocio"), "application/octet-stream"))])
    from app.models.models import Project
    project = db_session.query(Project).filter_by(company_id=cid).one()
    brief = Nivel1Service(llm=ScriptLLM([LEARNING]), db=db_session).board_brief(project)
    assert "Lo que ADÁN aprendió de la documentación del cliente" in brief
    assert "Vacío: No hay estructura de costos" in brief and "Riesgo: Un solo horno" in brief


def test_degraded_model_keeps_the_documents_but_does_not_invent_learning(client, db_session):
    cid, headers = _setup(client, "docsdeg@example.com")
    _use(client, ScriptLLM([LEARNING], degraded="Kimi no disponible (HTTPStatusError)"))
    body = client.post(f"/api/v1/nivel1/{cid}/documents/upload", headers=headers,
                       files=[("files", ("plan.docx", _docx("Plan de negocio"), "application/octet-stream"))]).json()
    assert body["documents"][0]["title"] == "plan.docx"
    assert body["discovery"]["done"] == 0 and body["learning"] == {}
    assert "no pude analizarlos" in body["reply"]


def test_only_unreadable_files_is_a_clear_error(client):
    cid, headers = _setup(client, "docsmal@example.com")
    _use(client, ScriptLLM([LEARNING]))
    resp = client.post(f"/api/v1/nivel1/{cid}/documents/upload", headers=headers,
                       files=[("files", ("foto.exe", b"MZ", "application/octet-stream"))])
    assert resp.status_code == 400 and "formato no admitido" in resp.json()["detail"]


def test_documents_of_another_company_are_off_limits(client):
    cid, _ = _setup(client, "duena@example.com")
    _, other = _setup(client, "intrusa@example.com")
    _use(client, ScriptLLM([LEARNING]))
    resp = client.post(f"/api/v1/nivel1/{cid}/documents/upload", headers=other,
                       files=[("files", ("plan.docx", _docx("Plan"), "application/octet-stream"))])
    assert resp.status_code == 404
